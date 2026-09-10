"""tests/test_art_manifest.py — ART-01: manifest id/version, state, validations.

Three things B-017's occurrence split does not cover on its own, which this
lot (docs/design/ART-1-artifacts.md, MAPA_REUTILIZACION's ART-01 row) asks
for:

* a version-chained manifest per artifact, appended to and never edited in
  place, carrying `call_id`;
* a `draft | generated | validated | reviewed | discarded` state machine
  where a failed format check is what keeps an artifact out of `validated`
  (and therefore out of `reviewed`, which only follows `validated`);
* format validations for DOCX (tables fit the page), XLSX (formulas marked
  for recalculation on open), PDF, JSON and image, wired into
  `artifact_store.collect()`/`persist()` so every freshly produced artifact
  gets a manifest without its producer asking for one.
"""
from __future__ import annotations

import os

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from core import database as db_mod, middleware
from core.database import Base
from src import artifact_identity as identity
from src import artifact_store
from src.contracts import ExecutionResult
from src.contracts.blob import ArtifactOccurrence


@pytest.fixture()
def own_database(tmp_path, monkeypatch):
    """Its own file, matching every other artifact-identity test in this
    suite: the shared in-memory database is not ground two artifact tests can
    stand on at once."""
    url = "sqlite:///" + (tmp_path / "manifest.db").as_posix()
    engine = create_engine(url, connect_args={"check_same_thread": False})
    monkeypatch.setattr(db_mod, "engine", engine)
    monkeypatch.setattr(db_mod, "SessionLocal",
                        sessionmaker(autocommit=False, autoflush=False, bind=engine))
    Base.metadata.create_all(bind=engine)
    monkeypatch.setattr(artifact_store, "ARTIFACT_STORE_DIR", str(tmp_path / "store"))
    monkeypatch.setattr(artifact_store, "ARTIFACT_RUNS_DIR", str(tmp_path / "runs"))
    yield engine
    engine.dispose()


def _occurrence(occ_id="occ_manifest_1", sha256="a" * 64):
    identity.ensure_blob(sha256=sha256, byte_size=10, filename=f"{sha256}.bin")
    value, _ = identity.ensure_occurrence(ArtifactOccurrence.parse({
        "id": occ_id, "blob_sha256": sha256, "kind": "text", "owner": "luis",
    }))
    return value


# ── version chain + state machine (identity.py) ─────────────────────────────

def test_a_manifest_cannot_attach_to_an_occurrence_that_does_not_exist(own_database):
    with pytest.raises(identity.ArtifactNotFound):
        identity.record_manifest_version("occ_ghost", state="generated")


def test_versions_chain_and_never_overwrite_an_earlier_one(own_database):
    occ = _occurrence()
    v1 = identity.record_manifest_version(occ.id, state="generated", call_id="call_01")
    v2 = identity.transition_state(occ.id, "validated", reason="checks passed")
    v3 = identity.transition_state(occ.id, "reviewed", reason="signed off")

    assert (v1["version"], v2["version"], v3["version"]) == (1, 2, 3)
    assert v2["previous_version_id"] == v1["id"]
    assert v3["previous_version_id"] == v2["id"]
    # call_id set at version 1 survives into later versions that did not
    # override it — the trace back to the tool call is not lost on transition.
    assert v2["call_id"] == v3["call_id"] == "call_01"

    history = identity.manifest_history(occ.id)
    assert [row["state"] for row in history] == ["generated", "validated", "reviewed"]
    assert identity.latest_manifest(occ.id)["state"] == "reviewed"


def test_an_invalid_transition_is_refused_and_nothing_is_written(own_database):
    occ = _occurrence()
    identity.record_manifest_version(occ.id, state="generated")
    with pytest.raises(ValueError, match="generated.*reviewed"):
        identity.transition_state(occ.id, "reviewed")
    # The refused jump left no trace: still exactly one version.
    assert len(identity.manifest_history(occ.id)) == 1


def test_discarded_is_terminal(own_database):
    occ = _occurrence()
    identity.record_manifest_version(occ.id, state="generated")
    identity.transition_state(occ.id, "discarded", reason="superseded")
    with pytest.raises(ValueError):
        identity.transition_state(occ.id, "generated")


def test_an_unknown_state_is_rejected(own_database):
    occ = _occurrence()
    with pytest.raises(ValueError):
        identity.record_manifest_version(occ.id, state="finished")


# ── format validations ──────────────────────────────────────────────────────

def test_docx_with_a_table_that_fits_the_page_validates(tmp_path):
    import docx
    path = str(tmp_path / "fits.docx")
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    document.save(path)
    result = identity.validate_artifact_bytes(path, filename="fits.docx")
    assert result["ok"] is True
    assert result["checks"][0]["name"] == "docx_tables_fit_page"


def test_docx_with_a_table_wider_than_the_page_fails_validation(tmp_path):
    import docx
    from docx.shared import Inches
    path = str(tmp_path / "overflows.docx")
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    for column in table.columns:
        column.width = Inches(3)  # 3 columns * 3in = 9in, page is 8.5in wide
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(3)
    document.save(path)
    result = identity.validate_artifact_bytes(path, filename="overflows.docx")
    assert result["ok"] is False
    assert result["checks"][0]["name"] == "docx_tables_fit_page"
    assert "EMU" in result["checks"][0]["detail"]


def test_a_docx_that_will_not_open_fails_with_a_named_reason(tmp_path):
    path = str(tmp_path / "corrupt.docx")
    with open(path, "wb") as fh:
        fh.write(b"not actually a docx")
    result = identity.validate_artifact_bytes(path, filename="corrupt.docx")
    assert result["ok"] is False
    assert result["checks"][0]["name"] == "docx_opens"


def test_xlsx_with_formulas_marked_for_recalculation_validates(tmp_path):
    import openpyxl
    from openpyxl.workbook.properties import CalcProperties
    path = str(tmp_path / "marked.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = 1, 2, "=A1+A2"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
    result = identity.validate_artifact_bytes(path, filename="marked.xlsx")
    assert result["ok"] is True
    assert "fullCalcOnLoad" in result["checks"][0]["detail"]


def test_xlsx_with_unmarked_formulas_fails_validation(tmp_path):
    """QA-39's literal scenario: a workbook with formulas whose recalculation
    is not guaranteed on open."""
    import openpyxl
    from openpyxl.workbook.properties import CalcProperties
    path = str(tmp_path / "unmarked.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = 1, 2, "=A1+A2"
    wb.calculation = CalcProperties(fullCalcOnLoad=False)
    wb.save(path)
    result = identity.validate_artifact_bytes(path, filename="unmarked.xlsx")
    assert result["ok"] is False
    assert result["checks"][0]["name"] == "xlsx_formulas_recalculate"
    assert "recalculating" in result["checks"][0]["detail"]


def test_xlsx_with_no_formulas_at_all_needs_no_recalc_marker(tmp_path):
    import openpyxl
    path = str(tmp_path / "plain.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "just data"
    wb.save(path)
    result = identity.validate_artifact_bytes(path, filename="plain.xlsx")
    assert result["ok"] is True
    assert "no formulas" in result["checks"][0]["detail"]


def test_pdf_that_parses_validates(tmp_path):
    import pypdf
    path = str(tmp_path / "ok.pdf")
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with open(path, "wb") as fh:
        writer.write(fh)
    result = identity.validate_artifact_bytes(path, filename="ok.pdf")
    assert result["ok"] is True and result["checks"][0]["name"] == "pdf_parses"


def test_a_pdf_that_will_not_parse_fails(tmp_path):
    path = str(tmp_path / "broken.pdf")
    with open(path, "wb") as fh:
        fh.write(b"this is not a pdf at all")
    result = identity.validate_artifact_bytes(path, filename="broken.pdf")
    assert result["ok"] is False


def test_valid_json_parses(tmp_path):
    path = str(tmp_path / "ok.json")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write('{"a": 1}')
    result = identity.validate_artifact_bytes(path, filename="ok.json")
    assert result["ok"] is True


def test_invalid_json_fails_to_parse(tmp_path):
    path = str(tmp_path / "broken.json")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write('{"a": ')
    result = identity.validate_artifact_bytes(path, filename="broken.json")
    assert result["ok"] is False and result["checks"][0]["name"] == "json_parses"


def test_a_decodable_image_validates(tmp_path):
    from PIL import Image
    path = str(tmp_path / "ok.png")
    Image.new("RGB", (4, 4), color="red").save(path)
    result = identity.validate_artifact_bytes(path, kind="image", filename="ok.png")
    assert result["ok"] is True


def test_an_undecodable_image_fails(tmp_path):
    path = str(tmp_path / "broken.png")
    with open(path, "wb") as fh:
        fh.write(b"\x89PNGnope, not a real image body")
    result = identity.validate_artifact_bytes(path, kind="image", filename="broken.png")
    assert result["ok"] is False


def test_a_kind_with_no_validator_reports_that_honestly_rather_than_guessing(tmp_path):
    path = str(tmp_path / "run.sh")
    with open(path, "w") as fh:
        fh.write("#!/bin/sh\necho hi\n")
    result = identity.validate_artifact_bytes(path, kind="code", filename="run.sh")
    assert result["ok"] is True
    assert result["checks"][0]["name"] == "no_validator"


# ── wired into collect()/persist() ──────────────────────────────────────────

def _result(names, **over):
    body = {"run_id": "run-manifest", "backend": "docker_workspace",
            "status": "completed", "exit_code": 0, "artifact_filenames": list(names)}
    body.update(over)
    return ExecutionResult.parse(body)


def test_a_freshly_collected_xlsx_is_auto_marked_and_ends_up_validated(own_database, tmp_path):
    """collect() fixes what it can (marks formulas for recalculation on the
    run's own file, before hashing) and persist() records that outcome as a
    `validated` manifest — QA-39's "recalculo" half."""
    import openpyxl
    from openpyxl.workbook.properties import CalcProperties

    src_dir = tmp_path / "run"
    src_dir.mkdir()
    path = src_dir / "report.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = 1, 2, "=A1+A2"
    wb.calculation = CalcProperties(fullCalcOnLoad=False)  # deliberately unmarked
    wb.save(str(path))

    collected = artifact_store.collect(
        _result(["report.xlsx"]), source_dir=str(src_dir),
        store_dir=artifact_store.ARTIFACT_STORE_DIR)
    assert len(collected.artifacts) == 1
    art = collected.artifacts[0]

    stored_path = artifact_store.path_of(art.filename)
    reopened = openpyxl.load_workbook(stored_path, data_only=False)
    assert reopened.calculation.fullCalcOnLoad is True, \
        "collect() should have marked the formula for recalculation before publishing"

    artifact_store.persist(collected.artifacts, call_id="call_xlsx_01")
    manifest = identity.latest_manifest(art.id)
    assert manifest["state"] == "validated"
    assert manifest["call_id"] == "call_xlsx_01"
    history = identity.manifest_history(art.id)
    assert [row["state"] for row in history] == ["generated", "validated"]


def test_a_docx_with_an_overflowing_table_stays_generated_not_validated(own_database, tmp_path):
    """QA-39's other half: a validation failure is what keeps the artifact at
    `generated` instead of `validated` (and therefore out of `reviewed`)."""
    import docx
    from docx.shared import Inches

    src_dir = tmp_path / "run"
    src_dir.mkdir()
    path = src_dir / "report.docx"
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    for column in table.columns:
        column.width = Inches(3)
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(3)
    document.save(str(path))

    collected = artifact_store.collect(
        _result(["report.docx"]), source_dir=str(src_dir),
        store_dir=artifact_store.ARTIFACT_STORE_DIR)
    art = collected.artifacts[0]
    artifact_store.persist(collected.artifacts, call_id="call_docx_01")

    manifest = identity.latest_manifest(art.id)
    assert manifest["state"] == "generated"
    assert manifest["validations"][0]["ok"] is False
    assert manifest["validations"][0]["name"] == "docx_tables_fit_page"


def test_persisting_the_same_occurrence_twice_appends_no_second_manifest(own_database, tmp_path):
    """`ensure_occurrence()` is idempotent on replay; the manifest bookkeeping
    that rides on `made` must be too, or a retried delivery would grow an
    unbounded chain of identical `generated` versions."""
    src_dir = tmp_path / "run"
    src_dir.mkdir()
    (src_dir / "note.txt").write_text("hello", encoding="utf-8")
    collected = artifact_store.collect(
        _result(["note.txt"]), source_dir=str(src_dir),
        store_dir=artifact_store.ARTIFACT_STORE_DIR)
    art = collected.artifacts[0]

    artifact_store.persist(collected.artifacts, call_id="call_a")
    artifact_store.persist(collected.artifacts, call_id="call_a")  # same event, retried

    assert len(identity.manifest_history(art.id)) == 1


# ── HTTP surface (routes/artifact_routes.py) ────────────────────────────────

def test_manifest_endpoint_returns_the_full_chain_and_review_requires_validated(
        own_database, tmp_path, monkeypatch):
    from src.owner_identity import DEFAULT_LOCAL_OWNER
    monkeypatch.setattr(middleware, "auth_disabled", lambda: True)
    from routes.artifact_routes import setup_artifact_routes
    app = FastAPI()
    app.include_router(setup_artifact_routes())
    http = TestClient(app)

    src_dir = tmp_path / "run"
    src_dir.mkdir()
    (src_dir / "note.txt").write_text("hello", encoding="utf-8")
    collected = artifact_store.collect(
        _result(["note.txt"]), source_dir=str(src_dir),
        store_dir=artifact_store.ARTIFACT_STORE_DIR, owner=DEFAULT_LOCAL_OWNER)
    art = collected.artifacts[0]
    artifact_store.persist(collected.artifacts, call_id="call_http")

    body = http.get(f"/api/artifacts/{art.id}/manifest").json()
    assert body["ok"] is True
    # note.txt has no format validator, so it stays `generated`, not `validated`.
    assert body["manifest"][-1]["state"] == "generated"

    refused = http.post(f"/api/artifacts/{art.id}/review")
    assert refused.status_code == 409

    identity.transition_state(art.id, "validated", reason="manually validated for the test")
    accepted = http.post(f"/api/artifacts/{art.id}/review")
    assert accepted.status_code == 200
    assert accepted.json()["manifest"]["state"] == "reviewed"
