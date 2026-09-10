"""QA-39 · Export defectuoso (docs/spec/v2/acceptance_scenarios.json).

Estimulo: DOCX abre pero tabla excede pagina; XLSX tiene formulas sin
recalcular.
Resultado exigido (literal): "Validacion visual/calculo advierte y
distingue archivo generado de revisado."

Requisitos: ART-03, ART-04.

Estado: verde, a nivel del manifiesto de artefactos (ART-01,
docs/design/ART-1-artifacts.md). `src/artifact_identity.py::validate_artifact_
bytes()` abre el DOCX/XLSX de verdad y responde exactamente el estimulo:
una tabla DOCX que excede el ancho de pagina falla `docx_tables_fit_page`
(python-docx, mide el ancho de columnas contra `section.page_width` menos
margenes); un XLSX con formulas sin `fullCalcOnLoad`/`forceFullCalc` falla
`xlsx_formulas_recalculate`. `src/artifact_store.py::collect()` ademas
corrige en el sitio las formulas de un XLSX recien producido (fija
`fullCalcOnLoad` antes de calcular su hash) y `persist()` adjunta esa
validacion al manifiesto del artefacto: si todo pasa, el estado pasa de
`generated` a `validated`; si algo falla, se queda en `generated` con el
fallo anotado — la distincion generado/revisado que pedia el escenario
(`reviewed` solo es alcanzable desde `validated`, nunca desde `generated`).

Lo que sigue sin cerrar (fuera del alcance de ficheros propios de este
lote): `routes/session_routes.py::export_session` sirve el .docx/.xlsx de
`src/chat_export_docx.py`/`chat_export.py` directamente, sin pasar el
archivo exportado por `artifact_store.collect()`/`persist()` ni por tanto
por esta validacion. Ver "Cambios necesarios en ficheros ajenos" en el
informe del lote para el punto de enganche exacto.
"""
import openpyxl
import pytest
from docx.shared import Inches
from openpyxl.workbook.properties import CalcProperties

from src import artifact_identity as identity

pytestmark = pytest.mark.qa_state("green")


def test_xlsx_with_unmarked_formulas_is_flagged_not_silently_trusted(tmp_path):
    """The literal stimulus: an XLSX whose formulas are not marked for
    recalculation on open."""
    path = str(tmp_path / "budget.xlsx")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["A2"], sheet["A3"] = 10, 20, "=A1+A2"
    workbook.calculation = CalcProperties(fullCalcOnLoad=False)
    workbook.save(path)

    result = identity.validate_artifact_bytes(path, filename="budget.xlsx")

    assert result["ok"] is False
    check = result["checks"][0]
    assert check["name"] == "xlsx_formulas_recalculate"
    assert check["ok"] is False


def test_docx_with_a_table_that_overflows_the_page_is_flagged(tmp_path):
    """The literal stimulus: a DOCX that opens fine but whose table is wider
    than the page allows."""
    import docx
    path = str(tmp_path / "report.docx")
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    for column in table.columns:
        column.width = Inches(3)  # 9in of table on an 8.5in page
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(3)
    document.save(path)

    result = identity.validate_artifact_bytes(path, filename="report.docx")

    assert result["ok"] is False
    check = result["checks"][0]
    assert check["name"] == "docx_tables_fit_page"
    assert check["ok"] is False


def test_a_failed_validation_distinguishes_generated_from_reviewed(own_database_for_qa39):
    """The literal required outcome: the artifact's manifest state tells apart
    "generated" (produced, unchecked-or-failed) from "reviewed" (a human
    signed off on something that actually validated). `reviewed` must never
    be reachable straight from a `generated` artifact whose validation
    failed — that is the state machine's own enforcement, not this test's."""
    identity.ensure_blob(sha256="c" * 64, byte_size=10, filename=("c" * 64) + ".xlsx")
    from src.contracts.blob import ArtifactOccurrence
    occ, _ = identity.ensure_occurrence(ArtifactOccurrence.parse({
        "id": "occ_qa39_generated", "blob_sha256": "c" * 64, "kind": "dataset",
        "owner": "luis",
    }))
    identity.record_manifest_version(
        occ.id, state="generated",
        validations=[{"name": "xlsx_formulas_recalculate", "ok": False,
                      "detail": "formulas present without fullCalcOnLoad"}])

    assert identity.latest_manifest(occ.id)["state"] == "generated"
    with pytest.raises(ValueError, match="generated.*reviewed"):
        identity.transition_state(occ.id, "reviewed")


@pytest.fixture()
def own_database_for_qa39(tmp_path, monkeypatch):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    from core import database as db_mod
    from core.database import Base

    url = "sqlite:///" + (tmp_path / "qa39.db").as_posix()
    engine = create_engine(url, connect_args={"check_same_thread": False})
    monkeypatch.setattr(db_mod, "engine", engine)
    monkeypatch.setattr(db_mod, "SessionLocal",
                        sessionmaker(autocommit=False, autoflush=False, bind=engine))
    Base.metadata.create_all(bind=engine)
    yield engine
    engine.dispose()
