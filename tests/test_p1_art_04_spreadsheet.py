"""ART-04 — reliable spreadsheets: type-preserving CSV import, ID/date
guessing avoidance, formula-injection-safe CSV export, and XLSX range writes
that keep formulas live and mark the workbook for recalculation.
"""
from __future__ import annotations

import openpyxl
import pytest

from src import spreadsheet as sheet


# ---- classify_cell_text -----------------------------------------------

def test_leading_zero_id_kept_as_text_not_coerced_to_a_number():
    value, kind, warning = sheet.classify_cell_text("00123")
    assert (value, kind) == ("00123", "text")
    assert "identifier" in warning


def test_long_digit_run_kept_as_text_to_avoid_precision_loss():
    long_id = "123456789012345678"  # 18 digits — a float would mangle this
    value, kind, warning = sheet.classify_cell_text(long_id)
    assert (value, kind) == (long_id, "text")
    assert warning


def test_ambiguous_slashed_date_is_flagged_not_guessed():
    value, kind, warning = sheet.classify_cell_text("01/02/03")
    assert (value, kind) == ("01/02/03", "text")
    assert "ambiguous date" in warning


def test_iso_date_is_parsed_unambiguously():
    import datetime as _dt
    value, kind, warning = sheet.classify_cell_text("2024-03-07")
    assert kind == "date"
    assert value == _dt.date(2024, 3, 7)
    assert warning is None


def test_ordinary_numbers_still_parse_as_numbers():
    assert sheet.classify_cell_text("42") == (42, "int", None)
    v, k, w = sheet.classify_cell_text("3.5")
    assert k == "float" and v == 3.5 and w is None


def test_formula_text_is_classified_not_evaluated():
    value, kind, warning = sheet.classify_cell_text("=SUM(A1:A2)")
    assert (value, kind, warning) == ("=SUM(A1:A2)", "formula", None)


# ---- import_csv_rows ----------------------------------------------------

def test_import_csv_rows_preserves_types_and_collects_warnings():
    text = "id,amount,joined\n007,19.99,2023-11-01\n1500000000000000,3,05/06/07\n"
    result = sheet.import_csv_rows(text)
    header, row1, row2 = result["rows"]
    assert header == ["id", "amount", "joined"]
    assert row1[0] == "007"        # ID-shaped — stayed text
    assert row1[1] == 19.99
    assert str(row1[2]) == "2023-11-01"
    assert row2[0] == "1500000000000000"  # too many digits — stayed text
    assert row2[2] == "05/06/07"          # ambiguous — stayed text
    assert any("identifier" in w for w in result["warnings"])
    assert any("ambiguous date" in w for w in result["warnings"])


# ---- CSV export / formula-injection guard --------------------------------

@pytest.mark.parametrize("payload", ["=cmd|'/c calc'!A1", "+1+1", "-2+3", "@SUM(1,2)"])
def test_export_csv_neutralizes_formula_leading_characters_by_default(payload):
    out = sheet.export_csv_rows([[payload]])
    # csv.writer quotes a field that starts with a special leading char if
    # needed; check the neutralizing apostrophe survived into the field body.
    assert "'" + payload in out or out.strip() == "'" + payload


def test_export_csv_allow_formulas_opts_out_of_neutralization():
    import csv as _csv
    import io as _io
    out = sheet.export_csv_rows([["=SUM(1,2)"]], allow_formulas=True)
    assert next(_csv.reader(_io.StringIO(out))) == ["=SUM(1,2)"]


def test_export_csv_leaves_ordinary_text_untouched():
    out = sheet.export_csv_rows([["hello", "42"]])
    assert out.strip() == "hello,42"


# ---- XLSX preview: stale formula detection -------------------------------

def test_read_workbook_preview_flags_formula_with_no_cached_value(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 2
    ws["A2"] = 3
    ws["A3"] = "=A1+A2"   # openpyxl never computes formulas — no cached value
    wb.save(path)

    preview = sheet.read_workbook_preview(str(path))
    s = preview["sheets"][0]
    assert s["needs_recalculation"] is True
    assert "A3" in s["stale_formula_cells"]
    a3 = next(c for row in s["rows"] for c in row if c["ref"] == "A3")
    assert a3["formula"] == "=A1+A2"
    assert a3.get("stale") is True


def test_read_workbook_preview_reports_no_recalculation_needed_without_formulas(tmp_path):
    path = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    ws["A2"] = 42
    wb.save(path)

    preview = sheet.read_workbook_preview(str(path))
    s = preview["sheets"][0]
    assert s["needs_recalculation"] is False
    assert s["stale_formula_cells"] == []


# ---- write_range ----------------------------------------------------------

def test_write_range_preserves_formula_id_and_date_types(tmp_path):
    path = tmp_path / "target.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    wb.save(path)

    result = sheet.write_range(str(path), "Data", "A1", [
        ["00123", "2024-01-15", "=1+1"],
    ])
    assert result["marked_recalculate"] is True
    assert "A1" in result["id_cells_forced_text"]

    wb2 = openpyxl.load_workbook(path, data_only=False)
    ws2 = wb2["Data"]
    assert ws2["A1"].value == "00123"
    assert ws2["A1"].number_format == "@"
    import datetime as _dt
    assert ws2["B1"].value == _dt.datetime(2024, 1, 15)
    assert ws2["C1"].data_type == "f"
    assert ws2["C1"].value == "=1+1"
    assert wb2.calculation.fullCalcOnLoad is True


def test_write_range_rejects_unknown_sheet(tmp_path):
    path = tmp_path / "target.xlsx"
    openpyxl.Workbook().save(path)
    with pytest.raises(sheet.SpreadsheetError):
        sheet.write_range(str(path), "NoSuchSheet", "A1", [["x"]])
