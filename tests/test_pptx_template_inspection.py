"""Template structure of PowerPoint packages: masters, layouts, themes, fingerprints, and the comparison against a
reference template (preserved, theme changed, layout edited, slide size changed, .potx reference)."""

import shutil
import zipfile

import pytest

from src.deliverable_inspection import DeliverableInspectionError, compare_presentation_templates, inspect_deliverable

pptx = pytest.importorskip("pptx")
from pptx.util import Inches  # noqa: E402


def _rewrite(path, part, fn):
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == part:
                data = fn(data.decode("utf-8")).encode("utf-8")
            dst.writestr(info, data)
    tmp.replace(path)


@pytest.fixture
def decks(tmp_path):
    template = tmp_path / "template.pptx"
    pptx.Presentation().save(template)
    deck = pptx.Presentation(str(template))
    s1 = deck.slides.add_slide(deck.slide_layouts[0])
    s1.shapes.title.text = "Defensa del TFM"
    s1.notes_slide.notes_text_frame.text = "Saludo"
    s2 = deck.slides.add_slide(deck.slide_layouts[1])
    s2.shapes.title.text = "Resultados"
    preserved = tmp_path / "preserved.pptx"
    deck.save(preserved)
    return template, preserved


@pytest.mark.asyncio
async def test_template_facts_list_masters_layouts_theme_and_usage(decks):
    template, preserved = decks
    result = await inspect_deliverable(str(preserved))
    t = result["facts"]["template"]
    assert len(t["masters"]) == 1 and len(t["layouts"]) == 11 and len(t["themes"]) == 1
    assert t["themes"][0]["name"] == "Office Theme" and t["themes"][0]["fonts"]["major"] == "Calibri"
    assert set(t["themes"][0]["colors"]) >= {"dk1", "lt1", "accent1", "accent6"}
    assert t["slide_size"]["named"] == "4:3" and t["slide_size"]["width_in"] == 10.0
    title = next(l for l in t["layouts"] if l["name"] == "Title Slide")
    assert title["used_by_slides"] == [1] and {"type": "ctrTitle", "idx": None} in title["placeholders"]
    assert "Title Slide" not in t["unused_layouts"] and "Blank" in t["unused_layouts"]
    slides = result["facts"]["slides"]
    assert slides[0]["layout_name"] == "Title Slide" and slides[1]["layout_name"] == "Title and Content"
    assert slides[0]["master"] == t["masters"][0]["part"]
    assert any("Template facts" in x for x in result["limitations"])


@pytest.mark.asyncio
async def test_deck_made_from_the_template_preserves_it(decks):
    template, preserved = decks
    result = await inspect_deliverable(str(preserved), compare_with=str(template))
    cmp = result["template_comparison"]
    assert cmp["template_preserved"] is True and cmp["identical_template_fingerprint"] is True
    assert len(cmp["layouts"]["identical"]) == 11 and cmp["layouts"]["changed_same_name"] == []
    assert cmp["slides_on_layouts_outside_reference"] == [] and cmp["slide_size"]["same"]


@pytest.mark.asyncio
async def test_changed_theme_colors_and_fonts_are_reported(decks):
    template, preserved = decks
    with zipfile.ZipFile(preserved) as z:
        theme = next(n for n in z.namelist() if n.startswith("ppt/theme/"))
    _rewrite(preserved, theme, lambda x: x.replace('<a:accent1><a:srgbClr val="4F81BD"/>',
                                                   '<a:accent1><a:srgbClr val="C0392B"/>')
                                         .replace('<a:majorFont><a:latin typeface="Calibri"/>',
                                                  '<a:majorFont><a:latin typeface="Georgia"/>'))
    cmp = (await inspect_deliverable(str(preserved), compare_with=str(template)))["template_comparison"]
    assert cmp["template_preserved"] is False
    assert cmp["themes"]["changed_same_name"] == ["Office Theme"]
    assert cmp["theme_color_changes"] == {"accent1": {"reference": "4F81BD", "deck": "C0392B"}}
    assert cmp["theme_font_changes"] == {"major": {"reference": "Calibri", "deck": "Georgia"}}
    assert len(cmp["layouts"]["identical"]) == 11                      # layouts untouched


@pytest.mark.asyncio
async def test_edited_layout_and_slides_on_it_are_flagged(decks):
    template, preserved = decks
    _rewrite(preserved, "ppt/slideLayouts/slideLayout1.xml",
             lambda x: x.replace("<p:cSld ", '<p:cSld xmlns:x="urn:edit" x:edited="1" ', 1))
    cmp = (await inspect_deliverable(str(preserved), compare_with=str(template)))["template_comparison"]
    assert cmp["layouts"]["changed_same_name"] == ["Title Slide"]
    assert cmp["slides_on_layouts_outside_reference"] == [{"slide": 1, "layout": "Title Slide", "reason": "edited layout"}]
    assert cmp["template_preserved"] is False


@pytest.mark.asyncio
async def test_slide_size_change_and_potx_reference(decks, tmp_path):
    template, _ = decks
    wide = pptx.Presentation(str(template))
    wide.slide_width, wide.slide_height = Inches(13.333), Inches(7.5)
    wide.slides.add_slide(wide.slide_layouts[6])
    wide_path = tmp_path / "wide.pptx"
    wide.save(wide_path)
    potx = tmp_path / "template.potx"
    shutil.copyfile(template, potx)
    result = await inspect_deliverable(str(wide_path), compare_with=str(potx))
    cmp = result["template_comparison"]
    assert cmp["slide_size"]["same"] is False and cmp["slide_size"]["deck"]["named"] == "16:9"
    assert cmp["template_preserved"] is False
    assert (await inspect_deliverable(str(potx)))["kind"] == "presentation"


@pytest.mark.asyncio
async def test_compare_with_rejects_non_presentations(decks, tmp_path):
    template, preserved = decks
    csv = tmp_path / "x.csv"; csv.write_text("a,b\n1,2\n")
    with pytest.raises(DeliverableInspectionError) as exc:
        await inspect_deliverable(str(csv), compare_with=str(template))
    assert exc.value.code == "invalid_arguments"
    with pytest.raises(DeliverableInspectionError):
        await inspect_deliverable(str(preserved), compare_with=str(csv))


def test_compare_is_pure_and_handles_packages_without_masters():
    empty = {"template": {"masters": [], "layouts": [], "themes": [], "slide_size": None, "template_fingerprint": None},
             "slides": []}
    cmp = compare_presentation_templates(empty, empty)
    assert cmp["template_preserved"] is False and cmp["identical_template_fingerprint"] is False


def test_tool_schema_and_handler_accept_compare_with():
    from src.agent_tools import FUNCTION_TOOL_SCHEMAS
    schema = next(x['function'] for x in FUNCTION_TOOL_SCHEMAS if x['function']['name'] == 'inspect_deliverable')
    assert 'compare_with' in schema['parameters']['properties']
    assert schema['parameters']['required'] == ['path']


@pytest.mark.asyncio
async def test_normal_handler_passes_the_reference_template(decks, monkeypatch):
    template, preserved = decks
    import src.tool_execution as te
    monkeypatch.setattr(te, "_resolve_tool_path", lambda value: value, raising=False)
    from src.agent_tools.media_tools import InspectDeliverableTool
    import json
    tool = InspectDeliverableTool()
    result = await tool.execute(json.dumps({"path": str(preserved), "compare_with": str(template)}), {})
    assert result["exit_code"] == 0, result
    assert result["deliverable"]["template_comparison"]["template_preserved"] is True
    bad = await tool.execute(json.dumps({"path": str(preserved), "compare_with": 3}), {})
    assert bad["exit_code"] == 1 and bad["error_code"] == "invalid_arguments"


def _repackage_with_absolute_targets(src, dst, moves):
    """Rewrite a package the way some SDK-based producers do: new relationship ids, absolute targets, parts moved."""
    import hashlib
    import posixpath
    import re
    with zipfile.ZipFile(src) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    out, id_maps = {}, {}
    for name, data in parts.items():
        if not name.endswith(".rels"):
            continue
        folder = name.rsplit("/_rels/", 1)[0] if "/_rels/" in name else ""
        source = (folder + "/" if folder else "") + name.rsplit("/", 1)[-1][:-5]
        mapping = id_maps.setdefault(source, {})

        def fix(m, folder=folder, source=source, mapping=mapping):
            rid, target = m.group(1), m.group(3)
            if 'TargetMode="External"' in m.group(0):
                return m.group(0)
            new_id = "R" + hashlib.sha1((source + rid).encode()).hexdigest()[:16]
            mapping[rid] = new_id
            full = posixpath.normpath(posixpath.join(folder, target))
            return m.group(0).replace(f'Id="{rid}"', f'Id="{new_id}"').replace(
                f'Target="{target}"', f'Target="/{moves.get(full, full)}"')

        text = re.sub(r'<Relationship Id="([^"]+)"( Type="[^"]+")? Target="([^"]+)"[^>]*/>', fix, data.decode("utf-8"))
        out[name] = text.encode("utf-8")
    for name, data in parts.items():
        if name.endswith(".rels"):
            continue
        if name.endswith(".xml") and name in id_maps:
            text = data.decode("utf-8")
            for old, new in id_maps[name].items():
                text = re.sub(rf'(r:[A-Za-z]+)="{old}"', rf'\1="{new}"', text)
            data = text.encode("utf-8")
        if name == "[Content_Types].xml":
            text = data.decode("utf-8")
            for old, new in moves.items():
                text = text.replace(f'PartName="/{old}"', f'PartName="/{new}"')
            data = text.encode("utf-8")
        out[moves.get(name, name)] = data
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in out.items():
            z.writestr(name, data)


@pytest.mark.asyncio
async def test_renumbered_ids_absolute_targets_and_moved_theme_keep_the_template(decks, tmp_path):
    template, preserved = decks
    sdk = tmp_path / "sdk.pptx"
    _repackage_with_absolute_targets(preserved, sdk, {"ppt/theme/theme1.xml": "ppt/slideMasters/theme/theme1.xml"})
    with zipfile.ZipFile(sdk) as z:
        rels = z.read("ppt/slideMasters/_rels/slideMaster1.xml.rels").decode()
    assert 'Target="/ppt/slideMasters/theme/theme1.xml"' in rels and 'Id="rId' not in rels
    result = await inspect_deliverable(str(sdk), compare_with=str(template))
    t = result["facts"]["template"]
    assert [x["part"] for x in t["themes"]] == ["ppt/slideMasters/theme/theme1.xml"]
    assert t["themes"][0]["fonts"]["major"] == "Calibri"
    assert [s["layout_name"] for s in result["facts"]["slides"]] == ["Title Slide", "Title and Content"]
    cmp = result["template_comparison"]
    assert cmp["template_preserved"] is True and cmp["identical_template_fingerprint"] is True
    assert cmp["theme_color_changes"] == {} and cmp["masters"]["identical"]


@pytest.mark.asyncio
async def test_edited_unnamed_master_is_matched_by_part(decks):
    template, preserved = decks
    _rewrite(preserved, "ppt/slideMasters/slideMaster1.xml",
             lambda x: x.replace("<p:cSld>", '<p:cSld xmlns:x="urn:edit" x:edited="1">', 1))
    cmp = (await inspect_deliverable(str(preserved), compare_with=str(template)))["template_comparison"]
    assert cmp["masters"]["changed_same_name"] == ["ppt/slideMasters/slideMaster1.xml"]
    assert cmp["masters"]["missing_from_deck"] == [] and cmp["masters"]["not_in_reference"] == []
    assert cmp["template_preserved"] is False


def _with_layout_image(src, dst, png):
    """Copy of the package whose Blank layout shows ``png`` (bytes) through relationship rId99."""
    pic = ('<p:pic><p:nvPicPr><p:cNvPr id="99" name="Logo"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
           '<p:blipFill><a:blip r:embed="rId99"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
           '<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
           '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
    layout = "ppt/slideLayouts/slideLayout7.xml"
    with zipfile.ZipFile(src) as z, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as out:
        for info in z.infolist():
            data = z.read(info.filename)
            if info.filename == layout:
                data = data.decode("utf-8").replace("</p:spTree>", pic + "</p:spTree>", 1).encode("utf-8")
            elif info.filename == "ppt/slideLayouts/_rels/slideLayout7.xml.rels":
                data = data.decode("utf-8").replace("</Relationships>", (
                    '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'relationships/image" Target="../media/logo.png"/></Relationships>')).encode("utf-8")
            out.writestr(info, data)
        out.writestr("ppt/media/logo.png", png)


@pytest.mark.asyncio
async def test_swapped_layout_image_changes_the_fingerprint(decks, tmp_path):
    template, _ = decks
    red, red_again, blue = tmp_path / "red.pptx", tmp_path / "red2.pptx", tmp_path / "blue.pptx"
    _with_layout_image(template, red, b"\x89PNG red")
    _with_layout_image(template, red_again, b"\x89PNG red")
    _with_layout_image(template, blue, b"\x89PNG blue")
    same = (await inspect_deliverable(str(red_again), compare_with=str(red)))["template_comparison"]
    assert same["identical_template_fingerprint"] is True
    cmp = (await inspect_deliverable(str(blue), compare_with=str(red)))["template_comparison"]
    assert cmp["layouts"]["changed_same_name"] == ["Blank"]
    assert len(cmp["layouts"]["identical"]) == 10


@pytest.mark.asyncio
async def test_aspect_ratio_of_a_non_standard_canvas(decks, tmp_path):
    template, _ = decks
    big = pptx.Presentation(str(template))
    big.slide_width, big.slide_height = Inches(20), Inches(11.25)
    path = tmp_path / "big.pptx"
    big.save(path)
    size = (await inspect_deliverable(str(path)))["facts"]["template"]["slide_size"]
    assert size["named"] is None and size["aspect"] == "16:9" and size["width_in"] == 20.0


@pytest.mark.asyncio
async def test_damaged_or_missing_template_parts_do_not_abort_the_inspection(decks, tmp_path):
    template, preserved = decks
    _rewrite(preserved, "ppt/slideLayouts/slideLayout11.xml", lambda x: x[: len(x) // 2])     # truncated XML
    result = await inspect_deliverable(str(preserved), compare_with=str(template))
    t = result["facts"]["template"]
    assert t["unreadable_parts"] == [{"part": "ppt/slideLayouts/slideLayout11.xml", "kind": "layout",
                                      "error": "invalid_package"}]
    assert len(t["layouts"]) == 10 and t["template_fingerprint"] is None
    assert [s["layout_name"] for s in result["facts"]["slides"]] == ["Title Slide", "Title and Content"]
    cmp = result["template_comparison"]
    assert cmp["template_preserved"] is False and cmp["unreadable_parts"]["deck"]
    no_theme = tmp_path / "no-theme.pptx"
    with zipfile.ZipFile(template) as src, zipfile.ZipFile(no_theme, "w") as dst:
        for info in src.infolist():
            if info.filename != "ppt/theme/theme1.xml":
                dst.writestr(info, src.read(info.filename))
    t = (await inspect_deliverable(str(no_theme)))["facts"]["template"]
    assert t["themes"] == [] and {"part": "ppt/theme/theme1.xml", "kind": "theme", "error": "missing_part"} in t["unreadable_parts"]


@pytest.mark.asyncio
async def test_external_link_targets_count_in_the_fingerprint(decks, tmp_path):
    template, _ = decks

    def with_link(dst, url):
        link = ('<p:sp><p:nvSpPr><p:cNvPr id="98" name="Link"><a:hlinkClick r:id="rId98"/></p:cNvPr><p:cNvSpPr/>'
                '<p:nvPr/></p:nvSpPr><p:spPr/></p:sp>')
        with zipfile.ZipFile(template) as z, zipfile.ZipFile(dst, "w") as out:
            for info in z.infolist():
                data = z.read(info.filename)
                if info.filename == "ppt/slideLayouts/slideLayout7.xml":
                    data = data.decode().replace("</p:spTree>", link + "</p:spTree>", 1).encode()
                elif info.filename == "ppt/slideLayouts/_rels/slideLayout7.xml.rels":
                    data = data.decode().replace("</Relationships>", (
                        '<Relationship Id="rId98" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                        f'relationships/hyperlink" Target="{url}" TargetMode="External"/></Relationships>')).encode()
                out.writestr(info, data)

    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    with_link(a, "https://example.org/a")
    with_link(b, "https://example.org/b")
    cmp = (await inspect_deliverable(str(b), compare_with=str(a)))["template_comparison"]
    assert cmp["layouts"]["changed_same_name"] == ["Blank"]


@pytest.mark.asyncio
async def test_broken_slide_layout_and_master_theme_links_are_not_preserved(decks):
    template, preserved = decks
    _rewrite(preserved, "ppt/slides/_rels/slide1.xml.rels",
             lambda x: x.replace("../slideLayouts/slideLayout1.xml", "../slideLayouts/missing.xml"))
    result = await inspect_deliverable(str(preserved), compare_with=str(template))
    assert result["facts"]["slides"][0]["layout_name"] is None
    assert result["facts"]["template"]["template_fingerprint"] is None
    cmp = result["template_comparison"]
    assert cmp["template_preserved"] is False and cmp["identical_template_fingerprint"] is False
    assert cmp["broken_links"]["deck"] == [{"from": "ppt/slides/slide1.xml", "kind": "slideLayout",
                                            "target": "ppt/slideLayouts/missing.xml", "error": "missing_part"}]

    import re
    _rewrite(preserved, "ppt/slides/_rels/slide1.xml.rels",
             lambda x: re.sub(r'<Relationship [^>]*slideLayout"[^>]*/>', "", x))
    cmp = (await inspect_deliverable(str(preserved), compare_with=str(template)))["template_comparison"]
    assert {"from": "ppt/slides/slide1.xml", "kind": "slideLayout", "target": None,
            "error": "no_relationship"} in cmp["broken_links"]["deck"]
    assert cmp["template_preserved"] is False


@pytest.mark.asyncio
async def test_master_without_theme_relationship_is_not_preserved(decks):
    template, preserved = decks
    import re
    _rewrite(preserved, "ppt/slideMasters/_rels/slideMaster1.xml.rels",
             lambda x: re.sub(r'<Relationship [^>]*relationships/theme"[^>]*/>', "", x))
    result = await inspect_deliverable(str(preserved), compare_with=str(template))
    t = result["facts"]["template"]
    assert t["themes"] == [] and t["template_fingerprint"] is None
    cmp = result["template_comparison"]
    assert cmp["template_preserved"] is False and cmp["identical_template_fingerprint"] is False
    assert {"from": "ppt/slideMasters/slideMaster1.xml", "kind": "theme", "target": None,
            "error": "no_relationship"} in cmp["broken_links"]["deck"]


def _facts(masters, layouts, themes):
    return {"template": {"slide_size": {"cx_emu": 1, "cy_emu": 1}, "masters": masters, "layouts": layouts,
                         "themes": themes, "template_fingerprint": "x", "unreadable_parts": [], "broken_links": []},
            "slides": []}


def test_same_parts_hung_differently_are_reported_as_foreign_arcs():
    themes = [{"part": "t1", "name": "A", "fingerprint": "T1", "colors": {}, "fonts": {}},
              {"part": "t2", "name": "B", "fingerprint": "T2", "colors": {}, "fonts": {}}]
    masters = [{"part": "m1", "name": "M1", "fingerprint": "M1", "theme": "t1", "theme_fingerprint": "T1"},
               {"part": "m2", "name": "M2", "fingerprint": "M2", "theme": "t2", "theme_fingerprint": "T2"}]
    layouts = [{"part": "l1", "name": "Title", "fingerprint": "L1", "master": "m1", "master_fingerprint": "M1"}]
    ref = _facts(masters, layouts, themes)
    moved = _facts(masters, [dict(layouts[0], master="m2", master_fingerprint="M2")], themes)
    cmp = compare_presentation_templates(ref, moved)
    assert cmp["layouts"]["identical"] == ["Title"] and cmp["masters"]["changed_same_name"] == []
    assert cmp["arcs_outside_reference"] == [{"from": "Title", "to": "M2"}]
    assert cmp["template_preserved"] is False
    assert compare_presentation_templates(ref, ref)["template_preserved"] is True


def test_theme_changes_are_reported_per_master_pair():
    ref_themes = [{"part": "t1", "name": "A", "fingerprint": "T1", "colors": {"accent1": "111111"}, "fonts": {"major": "Calibri"}},
                  {"part": "t2", "name": "B", "fingerprint": "T2", "colors": {"accent1": "222222"}, "fonts": {"major": "Arial"}}]
    deck_themes = [ref_themes[0], dict(ref_themes[1], fingerprint="T2x", colors={"accent1": "999999"})]
    masters = [{"part": "m1", "name": "Main", "fingerprint": "M1", "theme": "t1", "theme_fingerprint": "T1"},
               {"part": "m2", "name": "Closing", "fingerprint": "M2", "theme": "t2", "theme_fingerprint": "T2"}]
    deck_masters = [masters[0], dict(masters[1], theme_fingerprint="T2x")]
    cmp = compare_presentation_templates(_facts(masters, [], ref_themes), _facts(deck_masters, [], deck_themes))
    assert cmp["theme_changes"] == [{"master": "Closing", "reference_master": "Closing", "reference_theme": "t2",
                                     "deck_theme": "t2", "colors": {"accent1": {"reference": "222222", "deck": "999999"}},
                                     "fonts": {}}]
    assert cmp["theme_color_changes"] == {}            # the first master kept its theme
    assert cmp["themes"]["changed_same_name"] == ["B"] and cmp["template_preserved"] is False


@pytest.mark.asyncio
async def test_comparison_comes_before_the_slide_facts_in_the_serialised_result(decks):
    import json
    template, preserved = decks
    text = json.dumps(await inspect_deliverable(str(preserved), compare_with=str(template)), ensure_ascii=False)
    assert text.index('"template_comparison"') < text.index('"facts"') < text.index('"slides"')
